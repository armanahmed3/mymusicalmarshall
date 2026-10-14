import struct, os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(__file__))
from dump_all_database_tables import extract_table

# 1. MIXES (219)
mixes_raw = extract_table(219)
mixes_rows = []

singer_mapping = {
    2002: 'Luciano MM Mix by Luciano',
    2006: 'What Will Be by Stevie Malekuu',
    2007: "Don't Trust Dem by Stevie Malekuu",
    2008: 'War A Gwaan by Luciano',
    2010: 'Tennessee Whiskey by Wayne Armond',
    2011: 'Righteous People by Wayne Armond',
    2012: 'Wayne Armond Picks on Alton Ellis by Wayne Armond & Alton Ellis',
    3023: 'Test by Buju Banton',
    3025: 'Kush by Wayne Armond',
    3026: 'Chances Are by Yishka',
    3027: 'Hello Africa by Stevie Malekuu',
    3028: 'Groovy Reggae by Yishka',
    3029: 'Your Eyes by Teacha Barnes',
    3030: "Who's Loving You by Wayne Armond",
    3031: 'Heart Attack by Teacha Barnes',
    3032: 'Take A Chance by The Greaves Brothers',
}

for r in mixes_raw:
    fixed = r['fixed']
    mix_id = struct.unpack('<I', fixed[0:4])[0]
    mix_len_val = struct.unpack('<I', fixed[5:9])[0] / 100.0 if len(fixed) >= 9 else 0.0
    cols = r['var_cols']
    name = cols[0].strip() if len(cols) > 0 else ''
    raw_desc = cols[1].strip() if len(cols) > 1 else ''
    status = cols[2].strip() if len(cols) > 2 else ''
    audio = cols[3].strip() if len(cols) > 3 else ''
    cat = cols[4].strip() if len(cols) > 4 else ''
    genre = cols[5].strip() if len(cols) > 5 else ''
    
    dur_secs = int(round(mix_len_val * 60)) if mix_len_val > 0 else 240
    if dur_secs >= 3600:
        hrs = dur_secs // 3600
        mins = (dur_secs % 3600) // 60
        secs = dur_secs % 60
        dur_formatted = f"{hrs}:{mins:02d}:{secs:02d}"
    else:
        mins = dur_secs // 60
        secs = dur_secs % 60
        dur_formatted = f"{mins:02d}:{secs:02d}"
    
    if mix_id in singer_mapping:
        proper_desc = singer_mapping[mix_id]
    else:
        proper_desc = raw_desc if raw_desc else name

    mixes_rows.append({
        'mixId': mix_id,
        'mixName': name,
        'properDesc': proper_desc,
        'rawDesc': raw_desc,
        'category': cat,
        'genre': genre,
        'duration': dur_formatted,
        'lengthMinutes': f"{mix_len_val:.2f}",
        'status': status,
        'audioPath': audio
    })

# 2. USERS (212 & 230)
u212 = extract_table(212)
u230 = extract_table(230)
users_rows = []
seen_emails = set()

for u_list in [u230, u212]:
    for r in u_list:
        cols = r['var_cols']
        if not cols: continue
        email = cols[0].strip().lower()
        if not email or '@' not in email: continue
        if email in seen_emails: continue
        seen_emails.add(email)
        
        fname = cols[1].strip() if len(cols) > 1 else ''
        lname = cols[2].strip() if len(cols) > 2 else ''
        ref = cols[6].strip() if len(cols) > 6 else 'buju'
        status = cols[7].strip() if len(cols) > 7 else 'active'
        role = cols[9].strip() if len(cols) > 9 else 'user'
        utype = cols[10].strip() if len(cols) > 10 else 'regular'
        if 'admin' in email or 'orville' in email:
            role = 'admin'
            
        users_rows.append({
            'userId': len(users_rows) + 1,
            'email': email,
            'firstName': fname,
            'lastName': lname,
            'role': role,
            'referralCode': ref if ref else 'buju',
            'status': status if status else 'active',
            'userType': utype
        })

# 3. NOTICES (222)
notices_raw = extract_table(222)
notices_rows = []
for r in notices_raw:
    fixed = r['fixed']
    n_id = struct.unpack('<I', fixed[0:4])[0]
    cols = r['var_cols']
    img = cols[0].strip() if len(cols) > 0 else ''
    status = cols[1].strip() if len(cols) > 1 else 'active'
    loc = cols[2].strip() if len(cols) > 2 else ''
    title = cols[3].strip() if len(cols) > 3 else ''
    details = cols[4].strip() if len(cols) > 4 else ''
    notices_rows.append({
        'noticeId': n_id,
        'title': title,
        'location': loc,
        'details': details,
        'imagePath': img,
        'status': status
    })

# 4. PARAMETERS (225)
params_raw = extract_table(225)
p_cols = params_raw[0]['var_cols'] if params_raw else []
params_rows = [{
    'paramId': 1,
    'announcementNotice': p_cols[0].strip() if len(p_cols) > 0 else '',
    'featuredMix': p_cols[1].strip() if len(p_cols) > 1 else '',
    'defaultGenre': p_cols[2].strip() if len(p_cols) > 2 else ''
}]

# 5. REFERRAL CODES (229)
refs_raw = extract_table(229)
refs_rows = []
for idx, r in enumerate(refs_raw):
    code = r['var_cols'][0].strip() if r['var_cols'] else ''
    refs_rows.append({
        'id': idx + 1,
        'referralCode': code
    })

# 6. PLAYLISTS (228)
playlists_raw = extract_table(228)
playlists_rows = []
for idx, r in enumerate(playlists_raw):
    cols = r['var_cols']
    name = cols[0].strip() if len(cols) > 0 else ''
    status = cols[1].strip() if len(cols) > 1 else 'active'
    playlists_rows.append({
        'playlistId': idx + 1,
        'playlistName': name,
        'status': status
    })

# CREATE WORKBOOK
wb = openpyxl.Workbook()
wb.remove(wb.active) # remove default sheet

# Styling presets
header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
accent_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
data_font = Font(name="Calibri", size=10)
bold_font = Font(name="Calibri", size=10, bold=True)
thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

def format_sheet(ws, title, headers, data_list, key_order):
    ws.title = title
    ws.views.sheetView[0].showGridLines = True
    
    # Header row
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    
    ws.row_dimensions[1].height = 26
    
    # Data rows
    for row_idx, item in enumerate(data_list, start=2):
        row_vals = [item.get(k, '') for k in key_order]
        ws.append(row_vals)
        ws.row_dimensions[row_idx].height = 20
        for col_idx in range(1, len(row_vals) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in [1, 8, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

# Sheet 1: mixes
ws_mixes = wb.create_sheet(title="dbo.mixes")
format_sheet(
    ws_mixes,
    "dbo.mixes",
    [
        "mixId", 
        "mixName (Song Title)", 
        "Description (With Singer: This song is by...)", 
        "Raw DB mixDesc", 
        "mixCategory", 
        "Genre", 
        "Duration", 
        "mixLength (Min)", 
        "Status", 
        "audioPath"
    ],
    mixes_rows,
    ["mixId", "mixName", "properDesc", "rawDesc", "category", "genre", "duration", "lengthMinutes", "status", "audioPath"]
)

# Sheet 2: Users
ws_users = wb.create_sheet(title="dbo.Users")
format_sheet(
    ws_users,
    "dbo.Users",
    ["User ID", "Email Address", "First Name", "Last Name", "System Role", "Referral Code", "Account Status", "User Type"],
    users_rows,
    ["userId", "email", "firstName", "lastName", "role", "referralCode", "status", "userType"]
)

# Sheet 3: notices
ws_notices = wb.create_sheet(title="dbo.notices")
format_sheet(
    ws_notices,
    "dbo.notices",
    ["Notice ID", "Event Title", "Location", "Event Details", "Flyer Image Path", "Status"],
    notices_rows,
    ["noticeId", "title", "location", "details", "imagePath", "status"]
)

# Sheet 4: Parameters
ws_params = wb.create_sheet(title="dbo.Parameters")
format_sheet(
    ws_params,
    "dbo.Parameters",
    ["Param ID", "Announcement Notice", "Featured Mix", "Default Genre"],
    params_rows,
    ["paramId", "announcementNotice", "featuredMix", "defaultGenre"]
)

# Sheet 5: Referralcodes
ws_refs = wb.create_sheet(title="dbo.Referralcodes")
format_sheet(
    ws_refs,
    "dbo.Referralcodes",
    ["ID", "Referral Code"],
    refs_rows,
    ["id", "referralCode"]
)

# Sheet 6: Playlists
ws_playlists = wb.create_sheet(title="dbo.Playlists")
format_sheet(
    ws_playlists,
    "dbo.Playlists",
    ["Playlist ID", "Playlist Name", "Status"],
    playlists_rows,
    ["playlistId", "playlistName", "status"]
)

# Save to Desktop and project folders
desktop_file = r'C:\Users\92312\Desktop\MusicMarshall_Complete_Database.xlsx'
wb.save(desktop_file)

project_file = 'MusicMarshall_Complete_Database.xlsx'
wb.save(project_file)

public_file = 'public/MusicMarshall_Complete_Database.xlsx'
wb.save(public_file)

print(f"Excel export complete!")
print(f"  Desktop: {desktop_file}")
print(f"  Project: {project_file}")
print(f"  Public: {public_file}")
